"""Resume file validation, text extraction and section parsing. All inputs are untrusted."""
import io
import re
import zipfile
from datetime import date

import fitz  # PyMuPDF
from docx import Document

from app.core.errors import AppError
from app.schemas.analysis import ParsedResume
from app.services.skills import find_skills
from app.services.textutil import normalize_text

PDF_MAGIC = b"%PDF-"
ZIP_MAGIC = b"PK\x03\x04"
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_PDF_PAGES = 15
ALLOWED = {".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


def validate_upload(filename: str, data: bytes, max_bytes: int) -> str:
    """Returns the extension ('.pdf'|'.docx') or raises AppError. Checks size, extension AND magic bytes."""
    if not data:
        raise AppError(400, "empty_file", "The uploaded file is empty.")
    if len(data) > max_bytes:
        raise AppError(413, "file_too_large", f"File exceeds the {max_bytes // (1024 * 1024)} MB limit.")
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext not in ALLOWED:
        raise AppError(415, "unsupported_type", "Only PDF and DOCX resumes are supported.")
    if ext == ".pdf" and not data.startswith(PDF_MAGIC):
        raise AppError(415, "content_mismatch", "This file isn't a valid PDF despite its extension.")
    if ext == ".docx":
        if not data.startswith(ZIP_MAGIC):
            raise AppError(415, "content_mismatch", "This file isn't a valid DOCX despite its extension.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = z.namelist()
                if "word/document.xml" not in names:
                    raise AppError(415, "content_mismatch", "This file isn't a valid DOCX document.")
                if sum(i.file_size for i in z.infolist()) > MAX_DOCX_UNCOMPRESSED:
                    raise AppError(413, "file_too_large", "The document expands to an unsafe size.")
        except zipfile.BadZipFile:
            raise AppError(415, "content_mismatch", "This file isn't a valid DOCX document.")
    return ext


def extract_text(ext: str, data: bytes) -> str:
    try:
        if ext == ".pdf":
            with fitz.open(stream=data, filetype="pdf") as doc:
                if doc.needs_pass:
                    raise AppError(422, "encrypted_pdf", "This PDF is password-protected. Remove the password and re-upload.")
                if doc.page_count > MAX_PDF_PAGES:
                    raise AppError(422, "too_many_pages", f"Resumes over {MAX_PDF_PAGES} pages aren't supported.")
                return "\n".join(page.get_text("text") for page in doc)
        d = Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    except AppError:
        raise
    except Exception:  # corrupt/odd documents; never leak parser internals
        raise AppError(422, "unreadable_document", "We couldn't read this document. Try re-exporting it, or paste the text manually.")


# ---------- section parsing ----------
_HEADINGS = {
    "summary": r"(?:professional\s+)?summary|objective|career\s+objective|profile|about\s+me",
    "skills": r"(?:technical\s+|key\s+|core\s+)?skills(?:\s*(?:&|and)\s*tools)?|technologies|tech\s+stack|competencies",
    "experience": r"(?:work\s+|professional\s+|relevant\s+)?experience|employment(?:\s+history)?|internships?|work\s+history",
    "education": r"education(?:al\s+background)?|academics?|qualifications",
    "projects": r"(?:academic\s+|personal\s+|key\s+)?projects?",
    "certifications": r"certifications?(?:\s*(?:&|and)\s*courses)?|licenses?|courses",
    "achievements": r"achievements?|awards?(?:\s*(?:&|and)\s*honou?rs)?|honou?rs",
    "publications": r"publications?|research\s+papers?",
    "activities": r"extra-?\s?curricular(?:\s+activities)?|activities|positions?\s+of\s+responsibility|volunteer(?:ing)?(?:\s+experience)?|leadership",
}
_HEAD_RX = {k: re.compile(rf"^\s*(?:{v})\s*:?\s*$", re.I) for k, v in _HEADINGS.items()}

_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_PHONE = re.compile(r"\+?\d[\d\s().-]{8,16}\d")
_LINKEDIN = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[A-Za-z0-9_%-]+", re.I)
_GITHUB = re.compile(r"(?:https?://)?(?:www\.)?github\.com/[A-Za-z0-9_-]+", re.I)

MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
_MON = r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?"
_RANGE = re.compile(rf"(?:{_MON}\s+)?(\d{{4}})\s*(?:-|–|—|to)\s*(?:(?:{_MON}\s+)?(\d{{4}})|(present|current|now|ongoing))", re.I)

EDU_RANK = {"diploma": 1, "bachelor": 2, "master": 3, "phd": 4}
_EDU_PATTERNS = [
    ("phd", re.compile(r"\bph\.?\s?d\b|\bdoctorate\b", re.I)),
    ("master", re.compile(r"\bmaster'?s?\b|\bm\.?\s?tech\b|\bm\.?\s?sc\b|\bmca\b|\bm\.?\s?s\.?\b(?=\s+in|\s*\()|\bmba\b|\bm\.?\s?e\.?\b(?=\s+in)", re.I)),
    ("bachelor", re.compile(r"\bbachelor'?s?\b|\bb\.?\s?tech\b|\bb\.?\s?sc\b|\bbca\b|\bb\.?\s?e\.?\b(?=\s+in|\s*\()|\bb\.?\s?s\.?\b(?=\s+in)|\bbba\b|\bbcom\b", re.I)),
    ("diploma", re.compile(r"\bdiploma\b", re.I)),
]


def detect_education_level(text: str) -> str | None:
    best = None
    for level, rx in _EDU_PATTERNS:
        if rx.search(text) and (best is None or EDU_RANK[level] > EDU_RANK[best]):
            best = level
    return best


def split_sections(text: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for line in text.split("\n"):
        stripped = line.strip().strip("•-–—*|").strip()
        hit = next((k for k, rx in _HEAD_RX.items() if len(stripped) <= 45 and rx.match(stripped)), None)
        if hit:
            current = hit
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items() if "\n".join(v).strip()}


def estimate_years(experience_text: str, today: date | None = None) -> float | None:
    """Union of month-ranges found in the experience section, in years. Returns None when no dates are parsable."""
    today = today or date.today()
    spans = []
    for m in _RANGE.finditer(experience_text):
        sm, sy, em, ey, present = m.group(1), int(m.group(2)), m.group(3), m.group(4), m.group(5)
        start = sy * 12 + (MONTHS[sm.lower()[:3]] - 1 if sm else 0)
        if present:
            end = today.year * 12 + today.month - 1
        else:
            end = int(ey) * 12 + (MONTHS[em.lower()[:3]] - 1 if em else 11)
        if 1980 <= sy <= today.year + 1 and end >= start and end - start <= 12 * 50:
            spans.append((start, end))
    if not spans:
        return None
    spans.sort()
    merged = [list(spans[0])]
    for s, e in spans[1:]:
        if s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    return round(sum(e - s + 1 for s, e in merged) / 12, 1)


def parse_resume(text: str) -> ParsedResume:
    sections = split_sections(text)
    recognized = [k for k in sections if k != "header"]
    warnings: list[str] = []
    if len(text) < 200:
        warnings.append("Very little text was extracted. If your resume is an image or scanned PDF, paste the text manually.")
    if len(recognized) < 2:
        warnings.append("Few section headings were detected, so evidence checks will be less precise. Standard headings (Skills, Experience, Education, Projects) help.")
    contact = {}
    if m := _EMAIL.search(text):
        contact["email"] = m.group()
    for m in _PHONE.finditer(sections.get("header", text[:400])):
        if 10 <= len(re.sub(r"\D", "", m.group())) <= 15:
            contact["phone"] = m.group().strip()
            break
    if m := _LINKEDIN.search(text):
        contact["linkedin"] = m.group()
    if m := _GITHUB.search(text):
        contact["github"] = m.group()

    exp_text = sections.get("experience", "")
    years = estimate_years(exp_text) if exp_text else None
    note = ""
    if exp_text and years is None:
        note = "Experience section found but no date ranges could be parsed."
    elif not exp_text and len(recognized) >= 2:
        years, note = 0.0, "No experience section detected (internships/projects are not counted as years)."
    elif years is not None:
        note = "Estimated from date ranges in the experience section; internships listed there are included."

    edu_text = sections.get("education") or (text if len(recognized) < 2 else "")
    return ParsedResume(
        sections=sections, contact=contact, skills_found=find_skills(text), years_experience=years, years_note=note,
        education_level=detect_education_level(edu_text) if edu_text else None, warnings=warnings,
    )
